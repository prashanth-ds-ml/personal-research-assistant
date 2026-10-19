"""Phase 1a-0 eval tooling: error-analysis sheet, judge calibration, subject eval set (docs/phases/phase-1a0-measure-first.md)."""
import json
import sqlite3
from pathlib import Path

import pytest

from research_assistant import eval_tools as et

ROOT = Path(__file__).resolve().parents[1]

REPORT = """---
title: Toy
---

# Toy report

## Evidence status

| Section | Facts | Documents | Status |
|---|---|---|---|
| A | 2 | 1 | supported |

## Executive summary

The model uses 4B parameters [E1]. [single source] Two papers agree on loss masking [E2][E3]. [established]

The evidence is thin on GPU use. [inference]

## Section A

Training ran for one epoch [E2]. It used 35k pairs [E3]. [single source]

| Model | Metric | Value |
| :--- | :--- | :--- |
| Qwen3-1.7B | Accuracy | 28.8 [E1] |

## Sources

1. Paper one

## Evidence appendix

| ID | Claim | Verbatim quote | Source |
|---|---|---|---|
| E1 | The model has 4B parameters. | "Batch size \\| 4 \\|" | Card A |
| E2 | Loss is masked. | "loss on assistant tokens only" | Paper B |
| E3 | 35k pairs were used. | "approximately 35k instruction-formatted pairs" | Card A |
"""


def test_parse_evidence_handles_escaped_pipes():
    ev = et.parse_evidence(REPORT)
    assert set(ev) == {"E1", "E2", "E3"}
    assert ev["E1"]["quote"] == "Batch size | 4 |"
    assert ev["E1"]["source"] == "Card A"
    assert ev["E3"]["claim"] == "35k pairs were used."


def test_claims_split_on_labels_keep_citations_and_skip_meta_sections():
    claims = et.parse_claims(REPORT)
    texts = [c["text"] for c in claims]
    assert all("Evidence appendix" != c["section"] and "Sources" != c["section"] for c in claims)
    assert all(c["section"] != "Evidence status" for c in claims)
    first = claims[0]
    assert first["section"] == "Executive summary" and first["label"] == "single source" and first["cites"] == ["E1"]
    agree = next(c for c in claims if "loss masking" in c["text"])
    assert agree["label"] == "established" and agree["cites"] == ["E2", "E3"]
    assert any(c["label"] == "inference" and not c["cites"] for c in claims)       # uncited inference kept
    two = next(c for c in claims if "one epoch" in c["text"])
    assert "35k pairs" in two["text"] and two["cites"] == ["E2", "E3"]               # one label covers both sentences
    assert not any("[single source]" in t for t in texts)                             # tag stripped from the text


def test_table_rows_with_citations_are_claims():
    rows = [c for c in et.parse_claims(REPORT) if c["kind"] == "table"]
    assert len(rows) == 1 and rows[0]["cites"] == ["E1"] and "Qwen3-1.7B" in rows[0]["text"]


def test_sample_is_stratified_and_reproducible():
    pool = [{"report": r, "id": f"{r}-{i}"} for r, n in (("a", 40), ("b", 10), ("c", 5)) for i in range(n)]
    s1 = et.stratified_sample(pool, 20, seed=7)
    s2 = et.stratified_sample(pool, 20, seed=7)
    assert [x["id"] for x in s1] == [x["id"] for x in s2]
    assert len(s1) == 20
    counts = {r: sum(1 for x in s1 if x["report"] == r) for r in "abc"}
    assert counts["a"] > counts["b"] >= counts["c"] >= 2                              # proportional, every report present


def test_find_passage_matches_quote_despite_whitespace_and_escapes(tmp_path):
    db = tmp_path / "kb.sqlite"
    con = sqlite3.connect(db)
    con.executescript("CREATE TABLE docs (id INTEGER PRIMARY KEY, title TEXT, markdown TEXT);"
                      "CREATE TABLE chunks (id INTEGER PRIMARY KEY, doc_id INTEGER, text TEXT);")
    con.execute("INSERT INTO docs VALUES (1, 'Card A', 'x')")
    con.execute("INSERT INTO chunks VALUES (1, 1, 'Header\n| Batch size   | 4 |\n| Weight decay | 0.001 |')")
    con.commit()
    finder = et.PassageFinder(str(db))
    hit = finder.find("Batch size | 4 |", "Card A")
    assert hit and "Weight decay" in hit["passage"] and hit["doc"] == "Card A"
    assert finder.find("a quote that is nowhere", "Card A") is None


def test_P0_T01_stage_breakdown_sums_to_the_wrong_count():
    rows = [{"verdict": "ok", "stage": ""}, {"verdict": "wrong", "stage": "extraction subject"},
            {"verdict": "wrong", "stage": "extraction subject"}, {"verdict": "wrong", "stage": "writing stitch"},
            {"verdict": "", "stage": ""}]
    b = et.stage_breakdown(rows)
    assert b["labelled"] == 4 and b["wrong"] == 3 and b["ok"] == 1
    assert b["stages"] == {"extraction subject": 2, "writing stitch": 1}
    assert sum(b["stages"].values()) == b["wrong"]
    with pytest.raises(ValueError):
        et.stage_breakdown([{"verdict": "wrong", "stage": "made up stage"}])


def test_P0_T02_subject_cases_load_and_cover_the_known_bad_cases():
    cases = et.load_subject_cases(ROOT / "evals" / "subject_facts.jsonl")
    assert len(cases) >= 30
    for c in cases:
        assert c["passage"] and c["question"] and set(c["expected"]) >= {"system", "size", "stage"}
        assert c["expected"]["stage"] in et.STAGES
    origins = {c.get("origin", "") for c in cases}
    for known in ("E7", "E8", "E79", "E76", "E40"):
        assert any(known in o.split(",") for o in origins), known


def test_P0_T03_judge_agreement_on_a_toy_sheet():
    pairs = [("y", "y"), ("y", "n"), ("n", "n"), ("n", "n"), ("y", "y"), ("n", "y")]
    a = et.agreement(pairs)
    assert a["n"] == 6 and a["agree"] == 4
    assert a["rate"] == pytest.approx(4 / 6)
    assert a["confusion"] == {"y/y": 2, "y/n": 1, "n/y": 1, "n/n": 2}
    assert a["kappa"] == pytest.approx((4 / 6 - 0.5) / 0.5)
    assert et.agreement([])["rate"] is None


def test_reference_labels_use_the_human_label_first_then_agreed_labels():
    rows = [{"id": "1", "judge": "y", "claude": "y", "human": ""},     # judge and Claude agree: reference = y
            {"id": "2", "judge": "n", "claude": "y", "human": "y"},    # human settles it
            {"id": "3", "judge": "n", "claude": "y", "human": ""}]     # disputed, not settled: no reference
    ref = et.reference_labels(rows)
    assert ref == {"1": "y", "2": "y"}


def test_subject_jsonl_lines_are_valid_json():
    for line in (ROOT / "evals" / "subject_facts.jsonl").read_text(encoding="utf-8").splitlines():
        if line.strip():
            json.loads(line)


def test_labelling_workbook_round_trip(tmp_path):
    """The Excel workbook (guide sheet + claims sheet with dropdowns) gives back exactly what the labeller typed."""
    rows = [{"id": "C01", "report": "r", "section": "s", "kind": "prose", "label": "single source", "claim": "A claim.",
             "cites": "E1", "evidence": "E1 fact: ...\nquote: ..."},
            {"id": "C02", "report": "r", "section": "s", "kind": "table", "label": "", "claim": "| a | b |",
             "cites": "E2", "evidence": "E2 fact: ..."}]
    path = tmp_path / "claims.xlsx"
    et.write_label_workbook(rows, path, guide=[("How to label", ""), ("Step 1", "Read the claim")])
    import openpyxl
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["How to label", "Claims", "Dropdown options"]
    ws = wb["Claims"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["id", "claim"] and "human_verdict" in header and "human_stage" in header
    ws.cell(row=2, column=header.index("human_verdict") + 1, value="wrong")
    ws.cell(row=2, column=header.index("human_stage") + 1, value="extraction subject")
    ws.cell(row=3, column=header.index("human_verdict") + 1, value=" OK ")
    wb.save(path)
    labels = et.read_label_workbook(path)
    assert labels["C01"] == {"human_verdict": "wrong", "human_stage": "extraction subject", "human_note": ""}
    assert labels["C02"]["human_verdict"] == "ok"
    dvs = {str(dv.sqref).split(":")[0][0]: dv for dv in ws.data_validations.dataValidation}
    assert len(dvs) == 3                                                # verdict, stage and note dropdowns
    col = {h: openpyxl.utils.get_column_letter(header.index(h) + 1) for h in et.LABEL_COLUMNS}
    assert dvs[col["human_verdict"]].showErrorMessage and dvs[col["human_stage"]].showErrorMessage
    assert not dvs[col["human_note"]].showErrorMessage                 # note: suggestions, free text allowed
    assert all(dv.prompt for dv in dvs.values())                        # pop-up help on every label cell
    stages = [c.value for c in wb["Dropdown options"]["A"]]
    assert set(et.FAIL_STAGES) <= set(stages) and {"ok", "wrong"} <= set(stages)
    assert all(len(dv.prompt) <= 255 for dv in dvs.values())            # Excel's limit for input messages
