"""Evaluation tooling for Phase 1a-0 "Measure first" (docs/phases/phase-1a0-measure-first.md).

  parse_evidence(md) / parse_claims(md)   a saved report -> its evidence table and its claims
  stratified_sample(pool, n, seed)        a reproducible sample, proportional per report
  PassageFinder(kb).find(quote, title)    the library passage a quote came from (read-only)
  stage_breakdown(rows)                   error analysis tally: where the wrong claims first went wrong
  agreement(pairs) / reference_labels()   judge calibration against settled labels
  load_subject_cases(path)                the extraction subject eval set

Pure functions, no model calls: everything here must give the same answer every time.
"""
import json
import random
import re
import sqlite3
from collections import Counter
from pathlib import Path

from research_assistant.knowledge import split_front_matter

LABELS = ("established", "single source", "synthesis", "inference")
STAGES = ("pretraining", "continued pretraining", "fine-tuning", "post-training", "inference", "evaluation", "none", "unknown")
FAIL_STAGES = ("retrieval miss", "bad source", "extraction subject", "writing subject", "writing stitch", "label", "other")
META_SECTIONS = {"evidence status", "sources", "evidence appendix"}

_LABEL_TAG = re.compile(r"\[(" + "|".join(LABELS) + r")\]", re.I)
_CITE = re.compile(r"\[(E\d+)\]")
_SEP_ROW = re.compile(r"\|[\s:|-]+\|")


def _cells(row: str) -> list[str]:
    """Cells of one markdown table row; ``\\|`` inside a cell is a literal pipe, not a separator."""
    s = row.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [p.replace("\\|", "|").strip() for p in re.split(r"(?<!\\)\|", s)]


def _sections(md: str) -> list[tuple[str, str]]:
    """(heading, body) for every ``## `` section of the report body."""
    _, body = split_front_matter(md)
    out, head, buf = [], "", []
    for line in body.splitlines():
        if line.startswith("## "):
            if head:
                out.append((head, "\n".join(buf)))
            head, buf = line[3:].strip(), []
        elif head:
            buf.append(line)
    if head:
        out.append((head, "\n".join(buf)))
    return out


def parse_evidence(md: str) -> dict[str, dict]:
    """The evidence appendix as {id: {claim, quote, source}}."""
    table = next((b for h, b in _sections(md) if h.lower() == "evidence appendix"), "")
    out = {}
    for line in table.splitlines():
        if not re.match(r"\|\s*E\d+\s*\|", line):
            continue
        cells = _cells(line)
        if len(cells) < 4:
            continue
        quote = cells[2].strip()
        if len(quote) >= 2 and quote[0] == quote[-1] == '"':
            quote = quote[1:-1]
        out[cells[0]] = {"claim": cells[1], "quote": quote, "source": " | ".join(cells[3:]).strip()}
    return out


def _segments(text: str) -> list[tuple[str, str]]:
    """Split prose into (segment, label): each segment ends at an evidence-label tag. Text after the last tag has no
    label (older reports have no tags at all); it is split into sentences so a claim stays one sentence, not a paragraph."""
    out, start = [], 0
    for m in _LABEL_TAG.finditer(text):
        seg = text[start:m.start()].strip()
        if seg:
            out.append((seg, m.group(1).lower()))
        start = m.end()
    for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Z*])", text[start:].strip()):
        if sentence.strip():
            out.append((sentence.strip(), ""))
    return out


def parse_claims(md: str) -> list[dict]:
    """Every claim in the report body: prose segments (closed by an evidence label) and table rows that cite evidence.
    Meta sections (evidence status, sources, appendix) are skipped. A prose segment is kept if it cites evidence or
    carries a label; a cited table row is one claim."""
    claims = []
    for head, body in _sections(md):
        if head.lower() in META_SECTIONS:
            continue
        prose: list[str] = []

        def flush():
            for seg, label in _segments(" ".join(prose)):
                cites = list(dict.fromkeys(_CITE.findall(seg)))
                clean = re.sub(r"^[*\s]+", "", seg).strip()
                if (cites or label) and len(clean) > 15:
                    claims.append({"section": head, "kind": "prose", "text": clean, "label": label, "cites": cites})
            prose.clear()

        for line in body.splitlines():
            s = line.strip()
            if s.startswith("|"):
                flush()
                if _SEP_ROW.fullmatch(s):
                    continue
                cites = list(dict.fromkeys(_CITE.findall(s)))
                if cites:
                    claims.append({"section": head, "kind": "table", "text": s, "label": "", "cites": cites})
            elif not s or s in ("***", "---"):
                flush()
            else:
                prose.append(s)
        flush()
    return claims


def stratified_sample(pool: list[dict], n: int, seed: int = 0, key: str = "report", min_per: int = 2) -> list[dict]:
    """A reproducible sample of ``n`` items, proportional to each group's size (``key``), at least ``min_per`` from
    every group that has that many."""
    rng = random.Random(seed)
    groups: dict[str, list[dict]] = {}
    for item in pool:
        groups.setdefault(item[key], []).append(item)
    total = len(pool)
    quota = {g: min(len(items), max(min_per, round(n * len(items) / total))) for g, items in groups.items()}
    while sum(quota.values()) > n:                        # rounding overshoot: trim the largest group
        g = max(quota, key=lambda k: quota[k])
        quota[g] -= 1
    while sum(quota.values()) < min(n, total):            # undershoot: top up groups with spare items
        g = max((k for k in quota if quota[k] < len(groups[k])), key=lambda k: len(groups[k]) - quota[k])
        quota[g] += 1
    out = []
    for g in sorted(groups):
        out += rng.sample(groups[g], quota[g])
    return out


def _squash(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\\|", "|")).strip().lower()


class PassageFinder:
    """Find the library chunk a verbatim quote came from. Read-only: opens the index with ``mode=ro``."""

    def __init__(self, kb_path: str):
        self.con = sqlite3.connect(f"file:{Path(kb_path).as_posix()}?mode=ro", uri=True)
        self._chunks = [(t, _squash(x), x) for t, x in
                        self.con.execute("SELECT d.title, c.text FROM chunks c JOIN docs d ON d.id = c.doc_id")]

    def find(self, quote: str, title: str = "") -> dict | None:
        q = _squash(quote)
        if len(q) < 4:
            return None
        hits = [(t, raw) for t, sq, raw in self._chunks if q in sq]
        if not hits:
            return None
        best = next(((t, raw) for t, raw in hits if title and t.strip().lower() == title.strip().lower()), hits[0])
        return {"doc": best[0], "passage": best[1]}


def stage_breakdown(rows: list[dict], verdict: str = "verdict", stage: str = "stage") -> dict:
    """Tally labelled claims: ok / wrong, and the first failing stage of each wrong one. Unlabelled rows are skipped;
    an unknown stage name is an error (a typo must not silently become its own category)."""
    labelled = [r for r in rows if r.get(verdict, "").strip()]
    stages: Counter = Counter()
    for r in labelled:
        if r[verdict].strip().lower() == "wrong":
            s = r.get(stage, "").strip().lower()
            if s not in FAIL_STAGES:
                raise ValueError(f"unknown stage {s!r}; use one of {FAIL_STAGES}")
            stages[s] += 1
    wrong = sum(stages.values())
    return {"labelled": len(labelled), "wrong": wrong, "ok": len(labelled) - wrong, "stages": dict(stages)}


def agreement(pairs: list[tuple[str, str]]) -> dict:
    """Agreement between two labellers on yes/no labels: rate, Cohen's kappa, confusion counts (first/second)."""
    pairs = [(a.strip().lower(), b.strip().lower()) for a, b in pairs]
    n = len(pairs)
    conf = {f"{a}/{b}": sum(1 for x, y in pairs if (x, y) == (a, b)) for a in "yn" for b in "yn"}
    if not n:
        return {"n": 0, "agree": 0, "rate": None, "kappa": None, "confusion": conf}
    agree = conf["y/y"] + conf["n/n"]
    po = agree / n
    pa_y = (conf["y/y"] + conf["y/n"]) / n
    pb_y = (conf["y/y"] + conf["n/y"]) / n
    pe = pa_y * pb_y + (1 - pa_y) * (1 - pb_y)
    kappa = (po - pe) / (1 - pe) if pe < 1 else 1.0
    return {"n": n, "agree": agree, "rate": po, "kappa": kappa, "confusion": conf}


def reference_labels(rows: list[dict]) -> dict[str, str]:
    """The settled label per row: the human's label when given, else the label the judge and Claude agree on.
    A disputed row nobody settled has no reference (it is left out of calibration, not guessed)."""
    ref = {}
    for r in rows:
        human = r.get("human", "").strip().lower()
        judge, claude = r.get("judge", "").strip().lower(), r.get("claude", "").strip().lower()
        if human in ("y", "n"):
            ref[r["id"]] = human
        elif judge in ("y", "n") and judge == claude:
            ref[r["id"]] = judge
    return ref


LABEL_COLUMNS = ("human_verdict", "human_stage", "human_note")
WORKBOOK_COLUMNS = ("id", "claim", "label", "evidence", "section", "report") + LABEL_COLUMNS
OPTIONS_SHEET = "Dropdown options"

# (value, meaning) per label column; verdict and stage are strict lists, notes are suggestions (free text allowed)
VERDICT_OPTIONS = (("ok", "The sentence honestly says what its source says, about the right thing"),
                   ("wrong", "The sentence would mislead a reader: wrong model, size or stage, off-goal source, "
                             "overstated, or mislabelled"))
STAGE_OPTIONS = (("bad source", "1st check: the source does not fit the goal (speech model for a text goal, essay guide)"),
                 ("extraction subject", "2nd check: the FACT line already names the wrong model, size or stage, or "
                                        "turns the source into a general rule"),
                 ("writing subject", "3rd check: the fact is right, the claim drops or changes its subject or meaning"),
                 ("writing stitch", "4th check: facts from different sources merged into a claim none of them makes"),
                 ("label", "5th check: content fine, evidence label wrong ('established' needs 2+ documents)"),
                 ("retrieval miss", "The right source was never found (rare for a cited claim)"),
                 ("other", "None of the above; say what in human_note"))
NOTE_OPTIONS = (("wrong model", "The fact or claim names a different model than the source"),
                ("wrong size", "Parameter count does not match the source"),
                ("fine-tuning presented as pretraining", "A fine-tune's settings or hardware read as pretraining"),
                ("post-training presented as pretraining", "A post-training study's numbers read as pretraining"),
                ("speech source, goal is text", "ASR / speech translation used for a text-only goal"),
                ("off-goal source", "The document is not about the goal at all"),
                ("one project's settings as a general rule", "What one project did, stated as what you must do"),
                ("benchmark rule as a model requirement", "A test-harness rule stated as a property of the model"),
                ("overstated", "The source says less, or says it more cautiously"),
                ("merges different sources", "Two sources combined into something neither says"),
                ("'established' on one document", "Label needs two or more documents"),
                ("true but weak relevance", "Correct, but far from the goal (e.g. a 295B model in a 4B-8B report)"),
                ("subject named in the sentence before", "'The model' is fine once you read the previous sentence"),
                ("borderline", "Could go either way"))
LABEL_HELP = {"human_verdict": ("Verdict (required)", "ok = the sentence honestly says what its source says. "
                                "wrong = it would mislead you (wrong model/size/stage, off-goal, overstated, "
                                "mislabelled)."),
              "human_stage": ("Stage (only if wrong)", "First check that fails: 1 bad source, 2 extraction subject "
                              "(FACT wrong), 3 writing subject (claim changes the fact), 4 writing stitch, 5 label. "
                              "Else retrieval miss / other. Empty if ok."),
              "human_note": ("Note (optional)", "Pick a common reason or type your own words. Useful for wrong, "
                             "borderline or weakly relevant rows.")}


def write_label_workbook(rows: list[dict], path: str | Path, guide: list[tuple[str, str]]):
    """An Excel labelling workbook: a "How to label" sheet (``guide`` rows: heading, text), a "Claims" sheet with
    one row per claim, wrapped text, and dropdowns with pop-up help on the label columns (verdict and stage strict, so
    a label cannot be mistyped; note suggestions with free text allowed), and a sheet explaining every dropdown value."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    g = wb.active
    g.title = "How to label"
    g.column_dimensions["A"].width = 28
    g.column_dimensions["B"].width = 110
    for head, text in guide:
        g.append([head, text])
        r = g.max_row
        g.cell(r, 1).font = Font(bold=True, size=13 if not text else 11)
        for c in (1, 2):
            g.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Claims")
    ws.append(list(WORKBOOK_COLUMNS))
    for row in rows:
        ws.append([row.get(c, "") for c in WORKBOOK_COLUMNS])
    widths = {"id": 6, "claim": 60, "label": 13, "evidence": 90, "section": 22, "report": 18,
              "human_verdict": 14, "human_stage": 20, "human_note": 40}
    fill = PatternFill("solid", fgColor="FFF2CC")
    for i, col in enumerate(WORKBOOK_COLUMNS, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths[col]
        ws.cell(1, i).font = Font(bold=True)
        for r in range(1, ws.max_row + 1):
            ws.cell(r, i).alignment = Alignment(wrap_text=True, vertical="top")
            if col in LABEL_COLUMNS and r > 1:
                ws.cell(r, i).fill = fill
    ws.freeze_panes = "C2"

    opts = wb.create_sheet(OPTIONS_SHEET)                    # the dropdowns read their values from this sheet
    opts.column_dimensions["A"].width = 40
    opts.column_dimensions["B"].width = 100
    ranges, r = {}, 0
    for col, options in (("human_verdict", VERDICT_OPTIONS), ("human_stage", STAGE_OPTIONS),
                         ("human_note", NOTE_OPTIONS)):
        r += 1 if r == 0 else 2
        opts.cell(r, 1, col).font = Font(bold=True, size=12)
        opts.cell(r, 2, "meaning").font = Font(bold=True)
        ranges[col] = f"'{OPTIONS_SHEET}'!$A${r + 1}:$A${r + len(options)}"
        for value, meaning in options:
            r += 1
            opts.cell(r, 1, value)
            opts.cell(r, 2, meaning).alignment = Alignment(wrap_text=True, vertical="top")

    last = ws.max_row
    col_of = {c: ws.cell(1, i).column_letter for i, c in enumerate(WORKBOOK_COLUMNS, 1)}
    for col in LABEL_COLUMNS:
        dv = DataValidation(type="list", formula1=ranges[col], allow_blank=True)
        dv.promptTitle, dv.prompt = LABEL_HELP[col]
        dv.showInputMessage = True
        dv.showErrorMessage = col != "human_note"            # openpyxl's default is off; note = suggestions only
        if dv.showErrorMessage:
            dv.error, dv.errorTitle = f"Pick a value from the list (meanings on '{OPTIONS_SHEET}')", "Not a valid label"
        ws.add_data_validation(dv)
        dv.add(f"{col_of[col]}2:{col_of[col]}{last}")
    assert [v for v, _ in STAGE_OPTIONS] and {v for v, _ in STAGE_OPTIONS} == set(FAIL_STAGES)
    wb.active = 1
    wb.save(path)


def read_label_workbook(path: str | Path) -> dict[str, dict]:
    """The human labels typed into the workbook's "Claims" sheet, as {id: {human_verdict, human_stage, human_note}}
    (trimmed; verdict and stage lower-cased)."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)["Claims"]
    it = ws.iter_rows(values_only=True)
    header = [str(h or "") for h in next(it)]
    out = {}
    for values in it:
        rec = dict(zip(header, values))
        if not rec.get("id"):
            continue
        labels = {c: str(rec.get(c) or "").strip() for c in LABEL_COLUMNS}
        labels["human_verdict"] = labels["human_verdict"].lower()
        labels["human_stage"] = labels["human_stage"].lower()
        out[str(rec["id"])] = labels
    return out


def load_subject_cases(path: str | Path) -> list[dict]:
    """The extraction subject eval set: one JSON object per line with passage, question and expected fields."""
    cases = []
    for i, line in enumerate(Path(path).read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        case = json.loads(line)
        exp = case.get("expected", {})
        if exp.get("stage") not in STAGES:
            raise ValueError(f"line {i}: stage {exp.get('stage')!r} not in {STAGES}")
        cases.append(case)
    return cases
